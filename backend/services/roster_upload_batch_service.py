"""
Roster Upload Batch Service
"""
from datetime import date, datetime, time, timedelta
from io import BytesIO
import hashlib
from typing import Any, List, Optional
from uuid import UUID
from uuid import uuid4
from openpyxl import Workbook, load_workbook
from fastapi import HTTPException, status
from sqlalchemy.orm import Session
from sqlalchemy import and_, func
from models.roster_upload_batch import RosterUploadBatch
from models.roster_shift import RosterShift
from models.pegawai import Pegawai
from models.shift_kelompok import ShiftKelompok
from models.unit import Unit
from repositories.roster_upload_batch_repository import RosterUploadBatchRepository
from schemas.roster_upload_batch import (
    RosterUploadBatchCreate,
    RosterUploadBatchUpdate,
    RosterUploadBatchResponse,
)


VALID_UPLOAD_STATUS = {"UPLOADED", "VALIDATED", "IMPORTED", "FAILED"}

TEMPLATE_HEADERS = [
    "id_pegawai",
    "id_unit",
    "shift_kelompok_kode",
    "tanggal_shift",
    "jam_mulai",
    "jam_selesai",
    "nomor_sesi",
    "grace_telat_override_menit",
    "toleransi_pulang_cepat_override_menit",
    "catatan",
]


class RosterUploadBatchService:
    def __init__(self, db: Session):
        self.db = db
        self.repo = RosterUploadBatchRepository(db)

    def get_all(self, skip: int = 0, limit: int = 100) -> List[RosterUploadBatchResponse]:
        items = self.repo.get_all_with_relations(skip=skip, limit=limit)
        return [self._to_response(item) for item in items]

    def count_all(self) -> int:
        return self.repo.count()

    def get_by_id(self, batch_id: UUID) -> RosterUploadBatchResponse:
        item = self.repo.get_by_uuid_with_relations(batch_id)
        if not item:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Roster upload batch {batch_id} tidak ditemukan")
        return self._to_response(item)

    def create(self, payload: RosterUploadBatchCreate) -> RosterUploadBatchResponse:
        data = payload.model_dump()
        self._validate_payload(data)

        item = RosterUploadBatch(**data)
        self.db.add(item)
        self.db.commit()
        self.db.refresh(item)
        item = self.repo.get_by_uuid_with_relations(item.id)
        return self._to_response(item)

    def update(self, batch_id: UUID, payload: RosterUploadBatchUpdate) -> RosterUploadBatchResponse:
        item = self.repo.get_by_uuid_with_relations(batch_id)
        if not item:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Roster upload batch {batch_id} tidak ditemukan")

        data = payload.model_dump(exclude_unset=True)

        merged = {
            "period_start": data.get("period_start", item.period_start),
            "period_end": data.get("period_end", item.period_end),
            "upload_status": data.get("upload_status", item.upload_status),
            "total_rows": data.get("total_rows", item.total_rows),
            "valid_rows": data.get("valid_rows", item.valid_rows),
            "invalid_rows": data.get("invalid_rows", item.invalid_rows),
        }
        self._validate_payload(merged)

        for key, value in data.items():
            setattr(item, key, value)

        self.db.commit()
        self.db.refresh(item)
        item = self.repo.get_by_uuid_with_relations(item.id)
        return self._to_response(item)

    def delete(self, batch_id: UUID) -> None:
        item = self.repo.get_by_uuid(batch_id)
        if not item:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=f"Roster upload batch {batch_id} tidak ditemukan")

        self.db.delete(item)
        self.db.commit()

    def build_template_excel(self) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "roster_template"

        worksheet.append(TEMPLATE_HEADERS)
        worksheet.append([
            "P001",
            15,
            "JK_REGULER",
            "2026-02-23",
            "07:00",
            "14:00",
            1,
            10,
            0,
            "Shift pagi reguler",
        ])
        worksheet.append([
            "P002",
            15,
            "JK_SHIFT",
            "2026-02-23",
            "19:00",
            "07:00",
            1,
            10,
            0,
            "Shift malam lintas tanggal",
        ])

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output.getvalue()

    def import_from_excel(
        self,
        file_name: str,
        file_content: bytes,
        uploaded_by_pegawai: Optional[str] = None,
    ) -> dict[str, Any]:
        if not file_name.lower().endswith((".xlsx", ".xlsm")):
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="File harus berformat .xlsx atau .xlsm",
            )

        file_checksum = hashlib.sha256(file_content).hexdigest()

        existing_batch = (
            self.db.query(RosterUploadBatch)
            .filter(RosterUploadBatch.file_checksum == file_checksum)
            .order_by(RosterUploadBatch.created_at.desc())
            .first()
        )
        if existing_batch:
            existing_with_relations = self.repo.get_by_uuid_with_relations(existing_batch.id)
            existing_errors = []
            if existing_batch.error_summary and isinstance(existing_batch.error_summary, dict):
                existing_errors = existing_batch.error_summary.get("errors", [])

            return {
                "batch": self._to_response(existing_with_relations),
                "result": {
                    "total_rows": existing_batch.total_rows,
                    "valid_rows": existing_batch.valid_rows,
                    "invalid_rows": existing_batch.invalid_rows,
                    "errors": existing_errors,
                    "idempotent_reused": True,
                },
            }

        try:
            workbook = load_workbook(filename=BytesIO(file_content), data_only=True)
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Gagal membaca file Excel: {exc}",
            )

        worksheet = workbook.active
        if worksheet.max_row < 2:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="File Excel tidak memiliki data baris (minimal 1 baris data)",
            )

        header_values = [
            str(cell).strip() if cell is not None else ""
            for cell in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        ]

        missing_headers = [column for column in TEMPLATE_HEADERS if column not in header_values]
        if missing_headers:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Header wajib tidak lengkap: {', '.join(missing_headers)}",
            )

        column_index = {column: header_values.index(column) for column in TEMPLATE_HEADERS}

        available_pegawai = {
            str(row[0]): True
            for row in self.db.query(Pegawai.id_pegawai).all()
        }
        available_units = {
            int(row[0]): True
            for row in self.db.query(Unit.id_unit).all()
            if row[0] is not None
        }
        shift_kelompok_by_kode = {
            str(row[0]).upper(): int(row[1])
            for row in self.db.query(ShiftKelompok.kode, ShiftKelompok.id).all()
        }

        valid_payloads: list[dict[str, Any]] = []
        row_errors: list[dict[str, Any]] = []
        seen_keys: set[tuple[str, datetime, datetime, int]] = set()
        period_start: Optional[date] = None
        period_end: Optional[date] = None
        total_rows = 0

        for excel_row_number, row in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if row is None or all(value in (None, "") for value in row):
                continue

            total_rows += 1

            try:
                payload = self._parse_row(
                    row=row,
                    row_number=excel_row_number,
                    column_index=column_index,
                    available_pegawai=available_pegawai,
                    available_units=available_units,
                    shift_kelompok_by_kode=shift_kelompok_by_kode,
                )

                unique_key = (
                    payload["id_pegawai"],
                    payload["jam_mulai"],
                    payload["jam_selesai"],
                    payload["nomor_sesi"],
                )

                if unique_key in seen_keys:
                    raise ValueError("Duplikasi baris dalam file (pegawai + jam_mulai + jam_selesai + nomor_sesi)")
                seen_keys.add(unique_key)

                existing = self.db.query(RosterShift.id).filter(
                    and_(
                        RosterShift.id_pegawai == payload["id_pegawai"],
                        RosterShift.jam_mulai == payload["jam_mulai"],
                        RosterShift.jam_selesai == payload["jam_selesai"],
                        RosterShift.nomor_sesi == payload["nomor_sesi"],
                    )
                ).first()
                if existing:
                    raise ValueError("Data roster sudah ada di database")

                overlap_in_file = next(
                    (
                        existing_payload
                        for existing_payload in valid_payloads
                        if existing_payload["id_pegawai"] == payload["id_pegawai"]
                        and existing_payload["jam_mulai"] < payload["jam_selesai"]
                        and existing_payload["jam_selesai"] > payload["jam_mulai"]
                    ),
                    None,
                )
                if overlap_in_file:
                    raise ValueError("Overlap shift pada file untuk pegawai yang sama")

                overlap_existing = self.db.query(RosterShift.id).filter(
                    and_(
                        RosterShift.id_pegawai == payload["id_pegawai"],
                        RosterShift.jam_mulai < payload["jam_selesai"],
                        RosterShift.jam_selesai > payload["jam_mulai"],
                    )
                ).first()
                if overlap_existing:
                    raise ValueError("Overlap shift dengan data roster existing")

                valid_payloads.append(payload)

                if period_start is None or payload["tanggal_shift"] < period_start:
                    period_start = payload["tanggal_shift"]
                if period_end is None or payload["tanggal_shift"] > period_end:
                    period_end = payload["tanggal_shift"]

            except ValueError as exc:
                row_errors.append({"row": excel_row_number, "error": str(exc)})

        valid_rows = len(valid_payloads)
        invalid_rows = len(row_errors)

        if valid_rows == 0 and invalid_rows == 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Tidak ada data valid yang bisa diproses dari file Excel",
            )

        upload_status = "IMPORTED" if valid_rows > 0 else "FAILED"

        batch = RosterUploadBatch(
            file_name=file_name,
            file_checksum=file_checksum,
            period_start=period_start,
            period_end=period_end,
            uploaded_by_pegawai=uploaded_by_pegawai,
            upload_status=upload_status,
            total_rows=total_rows,
            valid_rows=valid_rows,
            invalid_rows=invalid_rows,
            error_summary={"errors": row_errors[:200]} if row_errors else None,
        )
        if self.db.bind and self.db.bind.dialect.name == "sqlite":
            batch.id = str(uuid4())
        self.db.add(batch)
        self.db.flush()

        for payload in valid_payloads:
            payload["upload_batch_id"] = batch.id
            roster_shift = RosterShift(**payload)
            if self.db.bind and self.db.bind.dialect.name == "sqlite":
                roster_shift.id = self._next_roster_shift_id()
            self.db.add(roster_shift)

        self.db.commit()
        self.db.refresh(batch)
        batch_with_relations = self.repo.get_by_uuid_with_relations(batch.id)

        return {
            "batch": self._to_response(batch_with_relations),
            "result": {
                "total_rows": total_rows,
                "valid_rows": valid_rows,
                "invalid_rows": invalid_rows,
                "errors": row_errors,
                "idempotent_reused": False,
            },
        }

    def _next_roster_shift_id(self) -> int:
        current_max = self.db.query(func.max(RosterShift.id)).scalar()
        return int(current_max or 0) + 1

    def _parse_row(
        self,
        row: tuple[Any, ...],
        row_number: int,
        column_index: dict[str, int],
        available_pegawai: dict[str, bool],
        available_units: dict[int, bool],
        shift_kelompok_by_kode: dict[str, int],
    ) -> dict[str, Any]:
        id_pegawai = self._to_str(row[column_index["id_pegawai"]])
        if not id_pegawai:
            raise ValueError("id_pegawai wajib diisi")
        if id_pegawai not in available_pegawai:
            raise ValueError(f"id_pegawai {id_pegawai} tidak ditemukan")

        tanggal_shift = self._parse_date_value(row[column_index["tanggal_shift"]], "tanggal_shift")

        jam_mulai = self._parse_datetime_value(
            tanggal_shift=tanggal_shift,
            value=row[column_index["jam_mulai"]],
            field_name="jam_mulai",
        )
        jam_selesai = self._parse_datetime_value(
            tanggal_shift=tanggal_shift,
            value=row[column_index["jam_selesai"]],
            field_name="jam_selesai",
        )

        if jam_selesai <= jam_mulai:
            jam_selesai = jam_selesai + timedelta(days=1)

        if jam_selesai <= jam_mulai:
            raise ValueError("jam_selesai harus lebih besar dari jam_mulai")

        id_unit_raw = row[column_index["id_unit"]]
        id_unit = None
        if id_unit_raw not in (None, ""):
            try:
                id_unit = int(id_unit_raw)
            except Exception:
                raise ValueError("id_unit harus berupa angka")
            if id_unit not in available_units:
                raise ValueError(f"id_unit {id_unit} tidak ditemukan")

        shift_kelompok_kode = self._to_str(row[column_index["shift_kelompok_kode"]])
        shift_kelompok_id = None
        if shift_kelompok_kode:
            shift_kelompok_id = shift_kelompok_by_kode.get(shift_kelompok_kode.upper())
            if not shift_kelompok_id:
                raise ValueError(f"shift_kelompok_kode {shift_kelompok_kode} tidak ditemukan")

        nomor_sesi_raw = row[column_index["nomor_sesi"]]
        nomor_sesi = 1 if nomor_sesi_raw in (None, "") else int(nomor_sesi_raw)
        if nomor_sesi <= 0:
            raise ValueError("nomor_sesi harus lebih besar dari 0")

        grace_raw = row[column_index["grace_telat_override_menit"]]
        grace_telat_override_menit = None if grace_raw in (None, "") else int(grace_raw)
        if grace_telat_override_menit is not None and grace_telat_override_menit < 0:
            raise ValueError("grace_telat_override_menit tidak boleh negatif")

        toleransi_raw = row[column_index["toleransi_pulang_cepat_override_menit"]]
        toleransi_pulang_cepat_override_menit = None if toleransi_raw in (None, "") else int(toleransi_raw)
        if toleransi_pulang_cepat_override_menit is not None and toleransi_pulang_cepat_override_menit < 0:
            raise ValueError("toleransi_pulang_cepat_override_menit tidak boleh negatif")

        catatan = self._to_str(row[column_index["catatan"]])

        return {
            "id_pegawai": id_pegawai,
            "shift_kelompok_id": shift_kelompok_id,
            "id_unit": id_unit,
            "tanggal_shift": tanggal_shift,
            "jam_mulai": jam_mulai,
            "jam_selesai": jam_selesai,
            "nomor_sesi": nomor_sesi,
            "grace_telat_override_menit": grace_telat_override_menit,
            "toleransi_pulang_cepat_override_menit": toleransi_pulang_cepat_override_menit,
            "status_roster": "AKTIF",
            "catatan": catatan,
            "source_row_number": row_number,
        }

    @staticmethod
    def _to_str(value: Any) -> str:
        if value is None:
            return ""
        return str(value).strip()

    @staticmethod
    def _parse_date_value(value: Any, field_name: str) -> date:
        if value is None or value == "":
            raise ValueError(f"{field_name} wajib diisi")
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        try:
            return datetime.fromisoformat(str(value)).date()
        except Exception:
            raise ValueError(f"{field_name} harus format tanggal valid (YYYY-MM-DD)")

    @staticmethod
    def _parse_datetime_value(tanggal_shift: date, value: Any, field_name: str) -> datetime:
        if value is None or value == "":
            raise ValueError(f"{field_name} wajib diisi")
        if isinstance(value, datetime):
            return value
        if isinstance(value, time):
            return datetime.combine(tanggal_shift, value)

        text = str(value).strip()
        try:
            parsed_datetime = datetime.fromisoformat(text)
            return parsed_datetime
        except Exception:
            pass

        for time_format in ["%H:%M:%S", "%H:%M"]:
            try:
                parsed_time = datetime.strptime(text, time_format).time()
                return datetime.combine(tanggal_shift, parsed_time)
            except Exception:
                continue

        raise ValueError(f"{field_name} harus format jam valid (HH:MM atau HH:MM:SS)")

    def _validate_payload(self, data: dict) -> None:
        period_start = data.get("period_start")
        period_end = data.get("period_end")
        if period_start and period_end and period_end < period_start:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="period_end tidak boleh lebih kecil dari period_start")

        upload_status = data.get("upload_status")
        if upload_status and upload_status not in VALID_UPLOAD_STATUS:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"upload_status tidak valid. Gunakan salah satu: {', '.join(sorted(VALID_UPLOAD_STATUS))}")

        for field_name in ["total_rows", "valid_rows", "invalid_rows"]:
            value = data.get(field_name)
            if value is not None and value < 0:
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"{field_name} tidak boleh negatif")

    @staticmethod
    def _to_response(item: RosterUploadBatch) -> RosterUploadBatchResponse:
        data = RosterUploadBatchResponse.model_validate(item).model_dump()
        data["uploaded_by_nama"] = item.uploader.nama if item.uploader else None
        return RosterUploadBatchResponse(**data)
