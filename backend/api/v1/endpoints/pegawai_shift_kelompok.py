"""
Pegawai Shift Kelompok Endpoints
CRUD assignment pegawai ke kelompok shift
"""
from io import BytesIO
from datetime import date as date_type
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from config.database import get_db
from schemas.pegawai_shift_kelompok import PegawaiShiftKelompokCreate, PegawaiShiftKelompokUpdate
from services.pegawai_shift_kelompok_service import PegawaiShiftKelompokService
from repositories.shift_kelompok_repository import ShiftKelompokRepository
from utils.dependencies import require_permission
from utils.permission_registry import PermissionKeys
from utils.response import success_response


router = APIRouter(prefix="/pegawai-shift-kelompok", tags=["Pegawai Shift Kelompok"])


@router.get("/", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_SHIFT_KELOMPOK_READ))])
def get_pegawai_shift_kelompok(
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
):
    service = PegawaiShiftKelompokService(db)
    items = service.get_all(skip=skip, limit=limit)
    total = service.count_all()
    return success_response(
        message="Pegawai shift kelompok retrieved successfully",
        data={"items": [item.model_dump() for item in items], "total": total, "skip": skip, "limit": limit},
    )


@router.post("/", response_model=dict, status_code=status.HTTP_201_CREATED, dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_SHIFT_KELOMPOK_CREATE))])
def create_pegawai_shift_kelompok(payload: PegawaiShiftKelompokCreate, db: Session = Depends(get_db)):
    service = PegawaiShiftKelompokService(db)
    created = service.create(payload)
    return success_response(message="Pegawai shift kelompok created successfully", data=created.model_dump())


@router.get("/template/download", dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_SHIFT_KELOMPOK_CREATE))])
def download_pegawai_shift_template():
    """Download Excel template for bulk pegawai shift kelompok import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template Shift Pegawai"

    headers = [
        "id_pegawai *",
        "shift_kelompok_kode *",
        "effective_start_date * (YYYY-MM-DD)",
        "effective_end_date (YYYY-MM-DD)",
        "is_default (TRUE/FALSE)",
        "catatan",
    ]
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Example row
    ws.append(["PEG001", "SHIFT-A", "2025-01-01", "", "TRUE", ""])

    col_widths = [18, 22, 32, 28, 22, 25]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template_shift_pegawai.xlsx"}
    )


@router.post("/import", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_SHIFT_KELOMPOK_CREATE))])
def import_pegawai_shift(
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """Bulk import pegawai shift kelompok from Excel file."""
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="File harus berformat .xlsx atau .xls")

    contents = file.file.read()
    wb = openpyxl.load_workbook(BytesIO(contents), data_only=True)
    ws = wb.active

    service = PegawaiShiftKelompokService(db)
    kelompok_repo = ShiftKelompokRepository(db)

    success_count = 0
    errors = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue

        id_pegawai = str(row[0]).strip() if row[0] is not None else ""
        kode_kelompok = str(row[1]).strip() if row[1] is not None else ""
        start_date_raw = str(row[2]).strip() if row[2] is not None else ""
        end_date_raw = str(row[3]).strip() if row[3] is not None else ""
        is_default_raw = str(row[4]).strip().upper() if row[4] is not None else "TRUE"
        catatan = str(row[5]).strip() if row[5] is not None else None

        # Validation
        if not id_pegawai:
            errors.append({"row": row_idx, "error": "id_pegawai wajib diisi"})
            continue
        if not kode_kelompok:
            errors.append({"row": row_idx, "id_pegawai": id_pegawai, "error": "shift_kelompok_kode wajib diisi"})
            continue
        if not start_date_raw:
            errors.append({"row": row_idx, "id_pegawai": id_pegawai, "error": "effective_start_date wajib diisi"})
            continue

        # Lookup shift kelompok
        kelompok = kelompok_repo.get_by_kode(kode_kelompok)
        if not kelompok:
            errors.append({"row": row_idx, "id_pegawai": id_pegawai, "error": f"Shift kelompok kode '{kode_kelompok}' tidak ditemukan"})
            continue

        # Parse dates
        try:
            effective_start_date = date_type.fromisoformat(start_date_raw)
        except ValueError:
            errors.append({"row": row_idx, "id_pegawai": id_pegawai, "error": f"Format effective_start_date tidak valid: {start_date_raw} (gunakan YYYY-MM-DD)"})
            continue

        effective_end_date = None
        if end_date_raw and end_date_raw.lower() not in ("none", "null", "-", ""):
            try:
                effective_end_date = date_type.fromisoformat(end_date_raw)
            except ValueError:
                errors.append({"row": row_idx, "id_pegawai": id_pegawai, "error": f"Format effective_end_date tidak valid: {end_date_raw}"})
                continue

        if not catatan or catatan.lower() in ("none", "null", "-", ""):
            catatan = None

        is_default = is_default_raw != "FALSE"

        try:
            payload = PegawaiShiftKelompokCreate(
                id_pegawai=id_pegawai,
                shift_kelompok_id=kelompok.id,
                effective_start_date=effective_start_date,
                effective_end_date=effective_end_date,
                is_default=is_default,
                catatan=catatan,
            )
            service.create(payload)
            success_count += 1
        except Exception as e:
            errors.append({"row": row_idx, "id_pegawai": id_pegawai, "error": str(e)})

    return success_response(
        data={"success": success_count, "errors": errors, "total": success_count + len(errors)},
        message=f"Import selesai: {success_count} berhasil, {len(errors)} gagal"
    )


@router.get("/{assignment_id}", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_SHIFT_KELOMPOK_READ))])
def get_pegawai_shift_kelompok_by_id(assignment_id: int, db: Session = Depends(get_db)):
    service = PegawaiShiftKelompokService(db)
    item = service.get_by_id(assignment_id)
    return success_response(message="Pegawai shift kelompok retrieved successfully", data=item.model_dump())


@router.put("/{assignment_id}", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_SHIFT_KELOMPOK_UPDATE))])
def update_pegawai_shift_kelompok(assignment_id: int, payload: PegawaiShiftKelompokUpdate, db: Session = Depends(get_db)):
    service = PegawaiShiftKelompokService(db)
    updated = service.update(assignment_id, payload)
    return success_response(message="Pegawai shift kelompok updated successfully", data=updated.model_dump())


@router.delete("/{assignment_id}", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_SHIFT_KELOMPOK_DELETE))])
def delete_pegawai_shift_kelompok(assignment_id: int, db: Session = Depends(get_db)):
    service = PegawaiShiftKelompokService(db)
    service.delete(assignment_id)
    return success_response(message="Pegawai shift kelompok deleted successfully")
