"""
User Endpoints - Admin Only
"""
from io import BytesIO
from typing import List
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from config.database import get_db
from schemas.user import UserCreate, UserUpdate, UserResponse, UserDetail
from services.user_service import UserService
from repositories.role_repository import RoleRepository
from utils.response import success_response
from utils.dependencies import require_permission, CommonQueryParams
from utils.permission_registry import PermissionKeys

router = APIRouter(prefix="/users", tags=["Users"])


@router.get("/", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.USERS_READ))])
def get_users(
    commons: CommonQueryParams = Depends(),
    db: Session = Depends(get_db)
):
    """
    Get all users (Admin only)
    
    Supports pagination with page and limit parameters
    """
    user_service = UserService(db)
    users = user_service.get_all(skip=commons.offset, limit=commons.limit)
    total = user_service.count_all()
    
    return success_response(
        data={"items": [user.model_dump() for user in users], "total": total, "skip": commons.offset, "limit": commons.limit},
        message="Users retrieved successfully"
    )


@router.post("/", response_model=dict, status_code=status.HTTP_201_CREATED, dependencies=[Depends(require_permission(PermissionKeys.USERS_CREATE))])
def create_user(
    user_data: UserCreate,
    db: Session = Depends(get_db)
):
    """
    Create new user (Admin only)
    
    - **username**: Unique username (min 3 characters)
    - **password**: Password (min 6 characters)
    - **id_pegawai**: Pegawai ID (optional)
    - **role_ids**: List of role IDs to assign
    - **is_active**: User status (default: true)
    """
    user_service = UserService(db)
    user = user_service.create(user_data)
    
    return success_response(
        data=user.model_dump(),
        message="User created successfully"
    )


@router.get("/template/download", dependencies=[Depends(require_permission(PermissionKeys.USERS_CREATE))])
def download_user_template():
    """Download Excel template for bulk user import."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Template User"

    headers = ["username *", "password (default: Absen@1234)", "id_pegawai", "role_names (e.g: user)", "is_active (TRUE/FALSE)"]
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Example row
    ws.append(["john.doe", "Absen@1234", "PEG001", "user", "TRUE"])

    # Column widths
    col_widths = [20, 30, 20, 25, 22]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template_user.xlsx"}
    )


@router.post("/import", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.USERS_CREATE))])
def import_users(
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """Bulk import users from Excel file."""
    if not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(status_code=400, detail="File harus berformat .xlsx atau .xls")

    contents = file.file.read()
    wb = openpyxl.load_workbook(BytesIO(contents), data_only=True)
    ws = wb.active

    user_service = UserService(db)
    role_repo = RoleRepository(db)

    success_count = 0
    errors = []

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue

        username = str(row[0]).strip() if row[0] is not None else ""
        password = str(row[1]).strip() if row[1] is not None else "Absen@1234"
        id_pegawai = str(row[2]).strip() if row[2] is not None else None
        role_names_raw = str(row[3]).strip() if row[3] is not None else ""
        is_active_raw = str(row[4]).strip().upper() if row[4] is not None else "TRUE"

        # Validation
        if not username:
            errors.append({"row": row_idx, "error": "username wajib diisi"})
            continue
        if len(username) < 3:
            errors.append({"row": row_idx, "username": username, "error": "username minimal 3 karakter"})
            continue
        if not password or len(password) < 6:
            password = "Absen@1234"

        # Normalize id_pegawai
        if not id_pegawai or id_pegawai.lower() in ("none", "null", "-", ""):
            id_pegawai = None

        # Parse role names → role IDs
        role_ids = []
        role_errors = []
        if role_names_raw:
            for rname in [n.strip() for n in role_names_raw.split(",") if n.strip()]:
                role = role_repo.get_by_name(rname)
                if role:
                    role_ids.append(role.id)
                else:
                    role_errors.append(rname)
        if role_errors:
            errors.append({"row": row_idx, "username": username, "error": f"Role tidak ditemukan: {', '.join(role_errors)}"})
            continue

        is_active = is_active_raw != "FALSE"

        try:
            user_data = UserCreate(
                username=username,
                password=password,
                id_pegawai=id_pegawai,
                role_ids=role_ids,
                is_active=is_active
            )
            user_service.create(user_data)
            success_count += 1
        except Exception as e:
            errors.append({"row": row_idx, "username": username, "error": str(e)})

    return success_response(
        data={"success": success_count, "errors": errors, "total": success_count + len(errors)},
        message=f"Import selesai: {success_count} berhasil, {len(errors)} gagal"
    )


@router.get("/{user_id}", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.USERS_READ))])
def get_user(
    user_id: int,
    db: Session = Depends(get_db)
):
    """Get user by ID (Admin only)"""
    user_service = UserService(db)
    user = user_service.get_by_id(user_id)
    
    return success_response(
        data=user.model_dump(),
        message="User retrieved successfully"
    )


@router.put("/{user_id}", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.USERS_UPDATE))])
def update_user(
    user_id: int,
    user_data: UserUpdate,
    db: Session = Depends(get_db)
):
    """
    Update user (Admin only)
    
    All fields are optional. Only provided fields will be updated.
    """
    user_service = UserService(db)
    user = user_service.update(user_id, user_data)
    
    return success_response(
        data=user.model_dump(),
        message="User updated successfully"
    )


@router.delete("/{user_id}", response_model=dict, status_code=status.HTTP_200_OK, dependencies=[Depends(require_permission(PermissionKeys.USERS_DELETE))])
def delete_user(
    user_id: int,
    db: Session = Depends(get_db)
):
    """Delete user (Admin only)"""
    user_service = UserService(db)
    user_service.delete(user_id)
    
    return success_response(
        data=None,
        message=f"User with id {user_id} deleted successfully"
    )
