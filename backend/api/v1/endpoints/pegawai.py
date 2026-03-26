"""
Pegawai Endpoints
Admin-only CRUD operations for employees
"""
from typing import List, Optional
from io import BytesIO
from datetime import date as date_type
from fastapi import APIRouter, Depends, HTTPException, Query, status, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from config.database import get_db
from schemas.pegawai import PegawaiCreate, PegawaiUpdate, PegawaiResponse
from services.pegawai_service import PegawaiService
from utils.response import success_response
from utils.dependencies import require_permission
from utils.permission_registry import PermissionKeys


router = APIRouter(prefix="/pegawai", tags=["Pegawai"])


@router.get("/", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_READ))])
async def get_pegawai(
    skip: int = 0,
    limit: int = 10,
    search: Optional[str] = None,
    id_unit: Optional[int] = Query(None, description="Filter pegawai berdasarkan unit"),
    db: Session = Depends(get_db)
):
    """
    Get all pegawai or search by name/NIP (Admin only)
    - **skip**: Number of items to skip (default: 0)
    - **limit**: Items per page (default: 10)
    - **search**: Search query for name or NIP (optional)
    - **id_unit**: Filter by unit ID (optional)
    """
    from models.pegawai import Pegawai as PegawaiModel
    from sqlalchemy import or_

    service = PegawaiService(db)

    if id_unit:
        # Filter langsung per unit via query
        q = db.query(PegawaiModel).filter(PegawaiModel.id_unit == id_unit)
        if search:
            q = q.filter(or_(
                PegawaiModel.nama.ilike(f"%{search}%"),
                PegawaiModel.nip.ilike(f"%{search}%"),
            ))
        total = q.count()
        pegawai_objs = q.order_by(PegawaiModel.nama).offset(skip).limit(limit).all()
        from schemas.pegawai import PegawaiResponse
        pegawai_list = [PegawaiResponse.model_validate(p) for p in pegawai_objs]
    elif search:
        total = service.count_search(search)
        pegawai_list = service.search(search, skip=skip, limit=limit)
    else:
        total = service.count_all()
        pegawai_list = service.get_all(skip=skip, limit=limit)

    return success_response(
        message="Pegawai retrieved successfully",
        data={
            "items": [p.model_dump() for p in pegawai_list],
            "total": total,
            "skip": skip,
            "limit": limit,
            "search": search,
            "id_unit": id_unit,
        }
    )

@router.post("/", response_model=dict, status_code=status.HTTP_201_CREATED, dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_CREATE))])
async def create_pegawai(
    id_pegawai: str = Form(...),
    nip: Optional[str] = Form(None),
    nama: Optional[str] = Form(None),
    jenis_kelamin: Optional[str] = Form(None),
    tempat_lahir: Optional[str] = Form(None),
    tanggal_lahir: Optional[str] = Form(None),
    alamat: Optional[str] = Form(None),
    id_unit: Optional[int] = Form(None),
    kepala_id_unit: Optional[int] = Form(None),
    status: Optional[str] = Form(None),
    nohp: Optional[str] = Form(None),
    foto: UploadFile = File(None),
    db: Session = Depends(get_db)
):
    """
    Create new pegawai (Admin only)
    
    - **id_pegawai**: Employee ID (required)
    - **nip**: NIP number (optional)
    - **nama**: Name (optional)
    - **jenis_kelamin**: Gender - L/P (optional)
    - **tempat_lahir**: Place of birth (optional)
    - **tanggal_lahir**: Date of birth (YYYY-MM-DD) (optional)
    - **alamat**: Address (optional)
    - **id_unit**: Unit ID (optional)
    - **kepala_id_unit**: Head Unit ID for approval routing (optional)
    - **status**: Status (optional)
    - **nohp**: Phone number / WA number (optional)
    - **foto**: Photo file (JPG/PNG) (optional)
    """
    service = PegawaiService(db)
    
    # Parse tanggal_lahir
    tgl_lahir = None
    if tanggal_lahir:
        try:
            tgl_lahir = date_type.fromisoformat(tanggal_lahir)
        except ValueError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Invalid date format. Use YYYY-MM-DD"
            )
    
    # Create PegawaiCreate schema
    pegawai_data = PegawaiCreate(
        id_pegawai=id_pegawai,
        nip=nip,
        nama=nama,
        jenis_kelamin=jenis_kelamin,
        tempat_lahir=tempat_lahir,
        tanggal_lahir=tgl_lahir,
        alamat=alamat,
        id_unit=id_unit,
        kepala_id_unit=kepala_id_unit,
        status=status,
        nohp=nohp,
    )
    
    pegawai = await service.create(pegawai_data, foto)
    
    return success_response(
        message="Pegawai created successfully",
        data=pegawai.model_dump()
    )


@router.get("/template/download", dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_CREATE))])
async def download_pegawai_template():
    """
    Download Excel template for bulk pegawai import.
    Returns .xlsx file with styled header row and one example row.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pegawai"

    columns = [
        ("id_pegawai *",               15),
        ("nip",                         18),
        ("nama",                         28),
        ("jenis_kelamin (L/P)",          20),
        ("tempat_lahir",                 20),
        ("tanggal_lahir (YYYY-MM-DD)",  25),
        ("alamat",                       35),
        ("id_unit",                      10),
        ("kepala_id_unit",               15),
        ("status (AKTIF/TIDAK_AKTIF)",   25),
        ("nohp",                         18),
    ]

    header_font  = Font(bold=True, color="FFFFFF")
    header_fill  = PatternFill("solid", fgColor="2563EB")
    header_align = Alignment(horizontal="center", vertical="center")

    for col, (header, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font  = header_font
        cell.fill  = header_fill
        cell.alignment = header_align
        ws.column_dimensions[cell.column_letter].width = width

    ws.row_dimensions[1].height = 22

    # Example data row
    example = [
        "EMP001", "199001012020011001", "Budi Santoso", "L",
        "Surabaya", "1990-01-01", "Jl. Merdeka No. 1", "1", "", "AKTIF", "08123456789",
    ]
    example_fill  = PatternFill("solid", fgColor="EFF6FF")
    for col, val in enumerate(example, 1):
        cell = ws.cell(row=2, column=col, value=val)
        cell.fill = example_fill

    # Freeze header row
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=template_import_pegawai.xlsx"},
    )


@router.post("/import", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_CREATE))])
async def import_pegawai(
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """
    Bulk import pegawai from Excel (.xlsx) file.
    Returns import summary: success count + per-row error details.
    """
    fname = (file.filename or "").lower()
    if not (fname.endswith(".xlsx") or fname.endswith(".xls")):
        raise HTTPException(status_code=400, detail="File harus berformat Excel (.xlsx atau .xls)")

    content = await file.read()
    try:
        wb = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="File Excel tidak dapat dibaca. Pastikan format valid.")

    ws  = wb.active
    rows = list(ws.iter_rows(values_only=True))

    if len(rows) < 2:
        raise HTTPException(status_code=400, detail="File tidak memiliki data (minimal 1 baris data setelah header).")

    service = PegawaiService(db)
    ok, errors = 0, []

    def cell_str(val):
        """Normalise cell value to clean string or None."""
        if val is None:
            return None
        s = str(val).strip()
        return s if s and s.lower() != "none" else None

    for row_idx, row in enumerate(rows[1:], start=2):
        # Skip completely empty rows
        if not any(row):
            continue

        id_pegawai = cell_str(row[0] if len(row) > 0 else None)
        if not id_pegawai:
            errors.append({"row": row_idx, "message": "id_pegawai wajib diisi"})
            continue

        try:
            # tanggal_lahir
            tanggal_lahir = None
            raw_tgl = cell_str(row[5] if len(row) > 5 else None)
            if raw_tgl:
                try:
                    parts = raw_tgl.split("-")
                    tanggal_lahir = date_type(int(parts[0]), int(parts[1]), int(parts[2]))
                except Exception:
                    errors.append({"row": row_idx, "id_pegawai": id_pegawai,
                                   "message": f"Format tanggal_lahir tidak valid: '{raw_tgl}'. Gunakan YYYY-MM-DD"})
                    continue

            # id_unit / kepala_id_unit — allow blank
            def to_int_or_none(v):
                s = cell_str(v)
                if not s:
                    return None
                try:
                    return int(float(s))
                except ValueError:
                    return None

            jk = cell_str(row[3] if len(row) > 3 else None)
            if jk:
                jk = jk.upper()
                if jk not in ("L", "P"):
                    errors.append({"row": row_idx, "id_pegawai": id_pegawai,
                                   "message": f"jenis_kelamin harus L atau P, dapat: '{jk}'"})
                    continue

            pegawai_data = PegawaiCreate(
                id_pegawai     = id_pegawai,
                nip            = cell_str(row[1] if len(row) > 1 else None),
                nama           = cell_str(row[2] if len(row) > 2 else None),
                jenis_kelamin  = jk,
                tempat_lahir   = cell_str(row[4] if len(row) > 4 else None),
                tanggal_lahir  = tanggal_lahir,
                alamat         = cell_str(row[6] if len(row) > 6 else None),
                id_unit        = to_int_or_none(row[7] if len(row) > 7 else None),
                kepala_id_unit = to_int_or_none(row[8] if len(row) > 8 else None),
                status         = (cell_str(row[9] if len(row) > 9 else None) or "AKTIF").upper(),
                nohp           = cell_str(row[10] if len(row) > 10 else None),
            )

            await service.create(pegawai_data, foto=None)
            ok += 1

        except HTTPException as exc:
            errors.append({"row": row_idx, "id_pegawai": id_pegawai, "message": exc.detail})
        except Exception as exc:
            errors.append({"row": row_idx, "id_pegawai": id_pegawai, "message": str(exc)})

    total = len(rows) - 1
    return success_response(
        message=f"Import selesai: {ok} berhasil, {len(errors)} gagal dari {total} baris data",
        data={"success": ok, "errors": errors, "total": total},
    )


@router.get("/{pegawai_id}", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_READ))])
async def get_pegawai_by_id(
    pegawai_id: str,
    db: Session = Depends(get_db)
):
    """
    Get pegawai by ID (Admin only)
    
    - **pegawai_id**: Employee ID
    """
    service = PegawaiService(db)
    pegawai = service.get_by_id(pegawai_id)
    
    return success_response(
        message="Pegawai retrieved successfully",
        data=pegawai.model_dump()
    )


@router.put("/{pegawai_id}", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_UPDATE))])
async def update_pegawai(
    pegawai_id: str,
    nip: Optional[str] = Form(None),
    nama: Optional[str] = Form(None),
    jenis_kelamin: Optional[str] = Form(None),
    tempat_lahir: Optional[str] = Form(None),
    tanggal_lahir: Optional[str] = Form(None),
    alamat: Optional[str] = Form(None),
    id_unit: Optional[int] = Form(None),
    kepala_id_unit: Optional[int] = Form(None),
    status: Optional[str] = Form(None),
    nohp: Optional[str] = Form(None),
    foto: UploadFile = File(None),
    db: Session = Depends(get_db)
):
    """
    Update pegawai (Admin only)
    
    - **pegawai_id**: Employee ID
    - **nip**: NIP number (optional)
    - **nama**: Name (optional)
    - **jenis_kelamin**: Gender - L/P (optional)
    - **tempat_lahir**: Place of birth (optional)
    - **tanggal_lahir**: Date of birth (YYYY-MM-DD) (optional)
    - **alamat**: Address (optional)
    - **id_unit**: Unit ID (optional)
    - **kepala_id_unit**: Head Unit ID for approval routing (optional)
    - **status**: Status (optional)
    - **nohp**: Phone number / WA number (optional)
    - **foto**: Photo file (JPG/PNG) (optional)
    """
    service = PegawaiService(db)
    
    # Parse tanggal_lahir
    tgl_lahir = None
    if tanggal_lahir:
        try:
            tgl_lahir = date_type.fromisoformat(tanggal_lahir)
        except ValueError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Invalid date format. Use YYYY-MM-DD"
            )
    
    # Build update data (only include fields that were provided)
    update_dict = {}
    if nip is not None:
        update_dict["nip"] = nip
    if nama is not None:
        update_dict["nama"] = nama
    if jenis_kelamin is not None:
        update_dict["jenis_kelamin"] = jenis_kelamin
    if tempat_lahir is not None:
        update_dict["tempat_lahir"] = tempat_lahir
    if tgl_lahir is not None:
        update_dict["tanggal_lahir"] = tgl_lahir
    if alamat is not None:
        update_dict["alamat"] = alamat
    if id_unit is not None:
        update_dict["id_unit"] = id_unit
    if kepala_id_unit is not None:
        # 0 = sentinel dari frontend: "hapus kepala_id_unit (set NULL)"
        update_dict["kepala_id_unit"] = kepala_id_unit if kepala_id_unit != 0 else None
    if status is not None:
        update_dict["status"] = status
    if nohp is not None:
        update_dict["nohp"] = nohp
    
    pegawai_data = PegawaiUpdate(**update_dict)
    
    pegawai = await service.update(pegawai_id, pegawai_data, foto)
    
    return success_response(
        message="Pegawai updated successfully",
        data=pegawai.model_dump()
    )


@router.delete("/{pegawai_id}", response_model=dict, dependencies=[Depends(require_permission(PermissionKeys.PEGAWAI_DELETE))])
async def delete_pegawai(
    pegawai_id: str,
    db: Session = Depends(get_db)
):
    """
    Delete pegawai (Admin only)
    
    - **pegawai_id**: Employee ID
    """
    service = PegawaiService(db)
    service.delete(pegawai_id)
    
    return success_response(
        message=f"Pegawai with id {pegawai_id} deleted successfully",
        data=None
    )
